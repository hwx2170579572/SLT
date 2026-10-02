"""Convert the hold35k method Markdown documents into Word (.docx).

The source documents are ordinary Markdown plus three local extensions:

  ``@eq <text>``   a display formula already written in Unicode (centred)
  ``@tex <text>``  a LaTeX source line (kept in Markdown, dropped from Word)
  ``$$ ... $$``    a display formula written in LaTeX (rendered to Unicode,
                   and re-listed verbatim in a trailing LaTeX appendix)

Inline ``$...$`` maths is rendered to Unicode as well, so the Word file reads
as text and the Markdown file keeps the paste-ready LaTeX for Word's equation
editor.

Chinese text needs an explicit East-Asian font, which python-docx does not set
through ``run.font.name``; ``_style_run`` writes the ``w:eastAsia`` attribute.

Usage:
    python -m tools.md_to_docx <file.md> [<file.md> ...]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LATIN = "Times New Roman"
EA_BODY = "宋体"
EA_HEAD = "黑体"
MONO = "Consolas"

SIZES = {"title": 18, "h1": 14, "h2": 12, "h3": 11, "body": 10.5, "small": 9}

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
MATH_SPAN = re.compile(r"\$\$(.+?)\$\$|\$([^$\n]+?)\$")
TABLE_ROW = re.compile(r"^\|.*\|$")
TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")

SUP = str.maketrans(
    "0123456789+-=()abcdefghijklmnoprstuvwxyz",
    "\u2070\u00b9\u00b2\u00b3\u2074\u2075\u2076\u2077\u2078\u2079\u207a\u207b\u207c"
    "\u207d\u207e\u1d43\u1d47\u1d9c\u1d48\u1d49\u1da0\u1d4d\u02b0\u2071\u02b2\u1d4f"
    "\u02e1\u1d50\u207f\u1d52\u1d56\u02b3\u02e2\u1d57\u1d58\u1d5b\u02b7\u02e3\u02b8\u1dbb")
SUB = str.maketrans("0123456789+-=()aehijklmnoprstuvx",
                    "\u2080\u2081\u2082\u2083\u2084\u2085\u2086\u2087\u2088\u2089"
                    "\u208a\u208b\u208c\u208d\u208e\u2090\u2091\u2095\u1d62\u2c7c\u2096"
                    "\u2097\u2098\u2099\u2092\u209a\u1d63\u209b\u209c\u1d64\u1d65\u2093")
SUP_KEYS = set("0123456789+-=()abcdefghijklmnoprstuvwxyz")
SUB_KEYS = set("0123456789+-=()aehijklmnoprstuvx")

GREEK = {
    "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
    "epsilon": "\u03b5", "varepsilon": "\u03b5", "zeta": "\u03b6", "eta": "\u03b7",
    "theta": "\u03b8", "iota": "\u03b9", "kappa": "\u03ba", "lambda": "\u03bb",
    "mu": "\u03bc", "nu": "\u03bd", "xi": "\u03be", "pi": "\u03c0", "rho": "\u03c1",
    "sigma": "\u03c3", "tau": "\u03c4", "upsilon": "\u03c5", "phi": "\u03c6",
    "chi": "\u03c7", "psi": "\u03c8", "omega": "\u03c9", "ell": "\u2113",
    "Gamma": "\u0393", "Delta": "\u0394", "Theta": "\u0398", "Lambda": "\u039b",
    "Sigma": "\u03a3", "Pi": "\u03a0", "Phi": "\u03a6", "Psi": "\u03a8",
    "Omega": "\u03a9", "nabla": "\u2207", "partial": "\u2202",
}

ACCENTS = {"hat": "\u0302", "bar": "\u0304", "tilde": "\u0303", "dot": "\u0307",
           "vec": "\u20d7", "overline": "\u0304"}
WRAPPERS = ("mathrm", "mathbf", "mathit", "text", "textrm", "textnormal", "mathsf",
            "mathtt", "mathcal", "mathbb", "mathfrak")

SYMBOLS = {
    "times": "\u00d7", "cdot": "\u00b7", "pm": "\u00b1", "mp": "\u2213",
    "le": "\u2264", "leq": "\u2264", "ge": "\u2265", "geq": "\u2265",
    "neq": "\u2260", "approx": "\u2248", "equiv": "\u2261", "sim": "\u223c",
    "propto": "\u221d", "ll": "\u226a", "gg": "\u226b",
    "in": "\u2208", "notin": "\u2209", "subset": "\u2282", "subseteq": "\u2286",
    "cup": "\u222a", "cap": "\u2229", "setminus": "\u2216", "emptyset": "\u2205",
    "wedge": "\u2227", "vee": "\u2228", "neg": "\u00ac", "forall": "\u2200",
    "exists": "\u2203", "infty": "\u221e", "partial": "\u2202", "nabla": "\u2207",
    "to": "\u2192", "rightarrow": "\u2192", "leftarrow": "\u2190",
    "Rightarrow": "\u21d2", "Leftarrow": "\u21d0", "mapsto": "\u21a6",
    "sum": "\u03a3", "prod": "\u03a0", "int": "\u222b", "sqrt": "\u221a",
    "quad": "\u2003", "qquad": "\u2003\u2003", ",": " ", ";": " ", "!": "",
    ":": " ", "ldots": "\u2026", "dots": "\u2026", "star": "\u2217",
    "circ": "\u2218", "odot": "\u2299", "oplus": "\u2295", "angle": "\u2220",
    "parallel": "\u2225", "perp": "\u22a5", "because": "\u2235",
    "ast": "\u2217", "Vert": "\u2016", "vert": "|", "mid": "|", "colon": ":",
    "langle": "\u27e8", "rangle": "\u27e9", "rightarrowtail": "\u21a3",
}

BLACKBOARD = {"R": "\u211d", "N": "\u2115", "Z": "\u2124", "Q": "\u211a", "C": "\u2102"}

DROPPED = {"left", "right", "big", "Big", "bigg", "Bigg", "bigl", "bigr", "Bigl",
           "Bigr", "displaystyle", "textstyle", "limits", "nolimits", "top",
           "middle", "ensuremath", "boldsymbol", "bm", "operatorname*"}


def _group(text: str, start: int) -> tuple[str, int]:
    """Return the content of the brace group opening at ``start``."""
    depth = 0
    for index in range(start, len(text)):
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:index], index + 1
    return text[start + 1:], len(text)


def _argument(text: str, start: int) -> tuple[str, int]:
    """Return the single argument (brace group, command or character) at ``start``."""
    index = start
    while index < len(text) and text[index].isspace():
        index += 1
    if index >= len(text):
        return "", index
    if text[index] == "{":
        return _group(text, index)
    if text[index] == "\\":
        match = re.match(r"\\[A-Za-z]+", text[index:])
        if match:
            return match.group(0), index + match.end()
    return text[index], index + 1


def _script(content: str, table, keys: set, marker: str) -> str:
    body = content.replace(" ", "")
    if body and all(char in keys for char in body):
        return body.translate(table)
    return marker + "(" + body + ")"


def _wrap_if_scripted(text: str, start: int, replacement: str) -> str:
    """Keep a marker group when a wrapper command directly follows ``^`` or ``_``."""
    if start > 0 and text[start - 1] in "^_":
        return "{" + replacement + "}"
    return replacement


def latex_to_text(latex: str, _final: bool = True) -> str:
    """Best-effort LaTeX to Unicode, readable enough inside Word."""
    text = latex
    # ``\_`` inside \text{...} is a literal underscore: hide it from the
    # subscript pass and restore it once the outermost call returns.
    text = text.replace("\\_", "\ue000")
    text = re.sub(r"\\tag\{[^}]*\}", "", text)
    text = re.sub(r"\\(?:begin|end)\{[^}]*\}", "", text)
    # Spacing / sizing commands go first: their arguments must not be parsed.
    text = text.replace("\\,", " ").replace("\\;", " ").replace("\\:", " ")
    text = text.replace("\\ ", " ").replace("\\!", "")
    for name in ("left", "right", "big", "Big", "bigg", "Bigg", "bigl", "bigr",
                 "Bigl", "Bigr", "displaystyle", "textstyle", "limits",
                 "nolimits", "middle", "ensuremath", "boldsymbol", "bm"):
        text = re.sub(r"\\" + name + r"(?![A-Za-z])", "", text)

    # Structural commands, rendered recursively on their arguments.
    while "\\frac" in text:
        start = text.index("\\frac")
        numerator, end = _argument(text, start + 5)
        denominator, end = _argument(text, end)
        top = latex_to_text(numerator, _final=False).strip()
        bottom = latex_to_text(denominator, _final=False).strip()
        simple = all(len(part) <= 3 and part.isalnum() for part in (top, bottom))
        body = f"{top}/{bottom}" if simple else f"({top})/({bottom})"
        text = text[:start] + body + text[end:]
    while "\\sqrt" in text:
        start = text.index("\\sqrt")
        content, end = _argument(text, start + 5)
        text = text[:start] + "\u221a(" + latex_to_text(content, _final=False).strip() \
            + ")" + text[end:]

    # Wrapper, accent and operator commands (innermost first, recursive).
    names = WRAPPERS + tuple(ACCENTS) + ("operatorname",)
    pattern = re.compile(r"\\(" + "|".join(names) + r")\*?(?![A-Za-z])")
    while True:
        match = pattern.search(text)
        if not match:
            break
        name = match.group(1)
        content, end = _argument(text, match.end())
        inner = latex_to_text(content, _final=False)
        if name in ACCENTS:
            replacement = inner + ACCENTS[name] if len(inner) <= 2 else inner
        elif name == "mathbb":
            replacement = "".join(BLACKBOARD.get(char, char) for char in inner)
        else:
            replacement = inner
        wrapped = _wrap_if_scripted(text, match.start(), replacement)
        text = text[:match.start()] + wrapped + text[end:]

    def command(match: re.Match) -> str:
        name = match.group(1)
        if name in GREEK:
            return GREEK[name]
        if name in SYMBOLS:
            return SYMBOLS[name]
        if name in DROPPED:
            return ""
        return name

    text = re.sub(r"\\([A-Za-z]+)", command, text)
    text = text.replace("\\{", "{").replace("\\}", "}").replace("\\|", "\u2016")
    text = text.replace("\\\\", " ")

    # Scripts run once commands are plain characters, then again after brace
    # removal so nested groups (e.g. ``_{s\in\{a,b\}}``) still resolve.
    def scripts(value: str) -> str:
        value = re.sub(r"\^\{([^{}]*)\}",
                       lambda m: _script(m.group(1), SUP, SUP_KEYS, "^"), value)
        value = re.sub(r"\^([A-Za-z0-9]+)",
                       lambda m: _script(m.group(1), SUP, SUP_KEYS, "^"), value)
        value = re.sub(r"_\{([^{}]*)\}",
                       lambda m: _script(m.group(1), SUB, SUB_KEYS, "_"), value)
        return re.sub(r"_([A-Za-z0-9]+)",
                      lambda m: _script(m.group(1), SUB, SUB_KEYS, "_"), value)

    text = scripts(text).replace("{", "").replace("}", "")
    text = scripts(text)
    text = re.sub(r"\s+", " ", text).strip()
    if _final:
        text = text.replace("", "_")
    return text


def _east_asian(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def _style_run(run, *, size: float, latin: str = LATIN, ea: str = EA_BODY,
               bold: bool = False, italic: bool = False,
               color: tuple[int, int, int] | None = None,
               mono: bool = False) -> None:
    name = MONO if mono else latin
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    _east_asian(run, MONO if mono else ea)


def _add_inline(paragraph, text: str, *, size: float, ea: str = EA_BODY,
                italic: bool = False, color: tuple[int, int, int] | None = None) -> None:
    """Render ``**bold**``, ``` `code` ``` and ``$math$`` spans."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            _add_math(paragraph, piece[2:-2], size=size, ea=ea, bold=True,
                      italic=italic, color=color)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            _style_run(paragraph.add_run(piece[1:-1]), size=size - 0.5, ea=ea,
                       italic=italic, color=color, mono=True)
        else:
            _add_math(paragraph, piece, size=size, ea=ea, italic=italic, color=color)


def _add_math(paragraph, text: str, *, size: float, ea: str = EA_BODY,
              bold: bool = False, italic: bool = False,
              color: tuple[int, int, int] | None = None) -> None:
    position = 0
    for match in MATH_SPAN.finditer(text):
        if match.start() > position:
            _style_run(paragraph.add_run(text[position:match.start()]), size=size,
                       ea=ea, bold=bold, italic=italic, color=color)
        latex = match.group(1) if match.group(1) is not None else match.group(2)
        _style_run(paragraph.add_run(latex_to_text(latex)), size=size, ea=ea,
                   bold=bold, italic=True, color=color)
        position = match.end()
    if position < len(text):
        _style_run(paragraph.add_run(text[position:]), size=size, ea=ea,
                   bold=bold, italic=italic, color=color)


def _shade(cell, hex_fill: str) -> None:
    prop = OxmlElement("w:shd")
    prop.set(qn("w:val"), "clear")
    prop.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(prop)


def _east_asian_style(style, name: str) -> None:
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def _setup(document: Document) -> None:
    for section in document.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.54)
        section.left_margin = section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(SIZES["body"])
    _east_asian_style(normal, EA_BODY)


def _heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    size = SIZES["h1"] if level == 1 else SIZES["h2"] if level == 2 else SIZES["h3"]
    _style_run(paragraph.add_run(text), size=size, ea=EA_HEAD, bold=True,
               color=(0, 0, 0))


def _table(document: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
    width = max(len(row) for row in cells)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, row in enumerate(cells):
        line = table.add_row().cells
        for column in range(width):
            text = row[column] if column < len(row) else ""
            paragraph = line[column].paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            _add_inline(paragraph, text, size=SIZES["small"],
                        ea=EA_HEAD if index == 0 else EA_BODY)
            for run in paragraph.runs:
                run.font.bold = index == 0 or run.font.bold
            if index == 0:
                _shade(line[column], "EDF2F7")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _code_block(document: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        _style_run(paragraph.add_run(line if line else " "),
                   size=SIZES["small"], mono=True)


def _equation(document: Document, rendered: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    _style_run(paragraph.add_run(rendered), size=SIZES["body"], italic=True)


def convert(source: Path, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    _setup(document)
    display_latex: list[str] = []

    index, first_heading_done = 0, False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("@tex"):
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index].rstrip())
                index += 1
            index += 1
            _code_block(document, block)
            continue

        if stripped.startswith("$$"):
            block = [stripped[2:]]
            if stripped.count("$$") < 2:
                index += 1
                while index < len(lines) and "$$" not in lines[index]:
                    block.append(lines[index].strip())
                    index += 1
                if index < len(lines):
                    block.append(lines[index].strip().split("$$")[0])
            else:
                block = [stripped[2:-2]]
            body = " ".join(part for part in block if part)
            display_latex.append(body)
            _equation(document, latex_to_text(body))
            index += 1
            continue

        if TABLE_ROW.match(stripped):
            block = []
            while index < len(lines) and TABLE_ROW.match(lines[index].strip()):
                block.append(lines[index].strip())
                index += 1
            _table(document, [row for row in block if not TABLE_SEP.match(row)])
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("@eq"):
            _equation(document, latex_to_text(stripped[3:].strip()))
            index += 1
            continue

        if stripped.startswith("#### "):
            _heading(document, stripped[5:], 3)
            index += 1
            continue
        if stripped.startswith("### "):
            _heading(document, stripped[4:], 2)
            index += 1
            continue
        if stripped.startswith("## "):
            _heading(document, stripped[3:], 1)
            index += 1
            continue
        if stripped.startswith("# "):
            title = stripped[2:]
            if first_heading_done:
                _heading(document, title, 1)
            else:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(10)
                _style_run(paragraph.add_run(title), size=SIZES["title"],
                           ea=EA_HEAD, bold=True)
                first_heading_done = True
            index += 1
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            _add_inline(paragraph, stripped.lstrip("> ").strip(), size=SIZES["small"],
                        italic=True, color=(0x44, 0x44, 0x44))
            index += 1
            continue

        bullet = re.match(r"^([-*]|\d+\.)\s+(.*)$", stripped)
        if bullet:
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet.group(1) in "-*" else "List Number")
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            _add_inline(paragraph, bullet.group(2), size=SIZES["body"])
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.35
        _add_inline(paragraph, stripped, size=SIZES["body"])
        index += 1

    if display_latex:
        _heading(document, "附录：本文公式的 LaTeX 源码", 1)
        note = document.add_paragraph()
        note.paragraph_format.space_after = Pt(6)
        _style_run(note.add_run("以下为正文各公式的 LaTeX 源码，可直接粘贴进 Word 公式编辑器或 LaTeX 排版。"),
                   size=SIZES["small"], italic=True, color=(0x44, 0x44, 0x44))
        for latex in display_latex:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.left_indent = Cm(0.6)
            _style_run(paragraph.add_run(latex), size=SIZES["small"], mono=True)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    for name in argv[1:]:
        source = Path(name)
        target = source.with_suffix(".docx")
        convert(source, target)
        print(f"{source.name} -> {target.name} ({target.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
